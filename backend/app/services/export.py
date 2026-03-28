"""Export service for generating Risk Log / CYA Package exports."""

import io
import tempfile
from datetime import datetime
from typing import Any

from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
    PageBreak,
)
from reportlab.lib.enums import TA_LEFT, TA_CENTER, TA_RIGHT
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

from app.schemas.analysis import RiskReport, SpecExtraction, DivisionData


class ExportService:
    """Service for generating PDF and Excel exports of risk reports."""

    def __init__(self):
        self.styles = getSampleStyleSheet()
        self._setup_custom_styles()

    def _setup_custom_styles(self):
        """Add custom paragraph styles."""
        self.styles.add(ParagraphStyle(
            name='Title2',
            parent=self.styles['Title'],
            fontSize=18,
            spaceAfter=12,
        ))
        self.styles.add(ParagraphStyle(
            name='SectionHeader',
            parent=self.styles['Heading2'],
            fontSize=14,
            spaceBefore=16,
            spaceAfter=8,
            textColor=colors.HexColor('#1a365d'),
        ))
        self.styles.add(ParagraphStyle(
            name='Subtitle',
            parent=self.styles['Normal'],
            fontSize=10,
            textColor=colors.grey,
            spaceAfter=20,
        ))
        self.styles.add(ParagraphStyle(
            name='RiskTitle',
            parent=self.styles['Heading3'],
            fontSize=11,
            spaceBefore=8,
            spaceAfter=4,
        ))
        self.styles.add(ParagraphStyle(
            name='RiskDescription',
            parent=self.styles['Normal'],
            fontSize=9,
            leftIndent=12,
        ))

    def _get_severity_color(self, severity: str) -> colors.Color:
        """Get color for severity level."""
        severity_colors = {
            'critical': colors.HexColor('#DC2626'),
            'high': colors.HexColor('#EA580C'),
            'medium': colors.HexColor('#CA8A04'),
            'low': colors.HexColor('#16A34A'),
        }
        return severity_colors.get(severity.lower(), colors.grey)

    def _get_gonogo_color(self, recommendation: str) -> colors.Color:
        """Get color for go/no-go recommendation."""
        gonogo_colors = {
            'proceed': colors.HexColor('#16A34A'),
            'proceed_with_contingency': colors.HexColor('#CA8A04'),
            'caution': colors.HexColor('#EA580C'),
            'do_not_bid': colors.HexColor('#DC2626'),
        }
        return gonogo_colors.get(recommendation.lower(), colors.grey)

    def _format_currency(self, amount: int | float | None) -> str:
        """Format number as currency."""
        if amount is None:
            return 'N/A'
        return f"${amount:,.0f}"

    def generate_pdf(
        self,
        risk_report: RiskReport,
        document_name: str,
        project_name: str | None = None,
        extraction: SpecExtraction | None = None,
    ) -> bytes:
        """Generate a PDF risk report / CYA package."""
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=letter,
            rightMargin=0.75*inch,
            leftMargin=0.75*inch,
            topMargin=0.75*inch,
            bottomMargin=0.75*inch,
        )

        elements = []

        # Title
        elements.append(Paragraph("BID RISK INTELLIGENCE REPORT", self.styles['Title2']))
        elements.append(Paragraph("CYA Package - Estimator Documentation", self.styles['Subtitle']))

        # Document Info
        doc_info = f"Document: {document_name}"
        if project_name:
            doc_info = f"Project: {project_name}<br/>{doc_info}"
        doc_info += f"<br/>Generated: {datetime.now().strftime('%B %d, %Y at %I:%M %p')}"
        elements.append(Paragraph(doc_info, self.styles['Normal']))
        elements.append(Spacer(1, 12))

        # Go/No-Go Recommendation Box
        if risk_report.go_no_go and isinstance(risk_report.go_no_go, dict):
            gonogo = risk_report.go_no_go
            rec = gonogo.get('recommendation', 'proceed').replace('_', ' ').upper()
            contingency = gonogo.get('contingency_percent', 0)
            reasoning = gonogo.get('reasoning', '')
            key_concerns = gonogo.get('key_concerns', [])

            elements.append(Paragraph("GO / NO-GO RECOMMENDATION", self.styles['SectionHeader']))

            # Recommendation table
            rec_data = [[
                Paragraph(f"<b>{rec}</b>", ParagraphStyle(
                    'RecStyle',
                    fontSize=14,
                    textColor=self._get_gonogo_color(gonogo.get('recommendation', 'proceed')),
                )),
                f"+{contingency}% Contingency" if contingency > 0 else "Standard Terms"
            ]]
            rec_table = Table(rec_data, colWidths=[3*inch, 3*inch])
            rec_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#F3F4F6')),
                ('TEXTCOLOR', (0, 0), (-1, -1), colors.black),
                ('ALIGN', (0, 0), (0, -1), 'LEFT'),
                ('ALIGN', (1, 0), (1, -1), 'RIGHT'),
                ('FONTSIZE', (0, 0), (-1, -1), 12),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 12),
                ('TOPPADDING', (0, 0), (-1, -1), 12),
                ('LEFTPADDING', (0, 0), (-1, -1), 12),
                ('RIGHTPADDING', (0, 0), (-1, -1), 12),
            ]))
            elements.append(rec_table)
            elements.append(Spacer(1, 6))

            if reasoning:
                elements.append(Paragraph(f"<i>{reasoning}</i>", self.styles['Normal']))
                elements.append(Spacer(1, 6))

            if key_concerns:
                elements.append(Paragraph("<b>Key Concerns:</b>", self.styles['Normal']))
                for concern in key_concerns[:3]:
                    elements.append(Paragraph(f"• {concern}", self.styles['RiskDescription']))
                elements.append(Spacer(1, 12))

        # Financial Exposure Summary
        if risk_report.financial_exposure and isinstance(risk_report.financial_exposure, dict):
            fin = risk_report.financial_exposure
            elements.append(Paragraph("FINANCIAL EXPOSURE SUMMARY", self.styles['SectionHeader']))

            fin_data = [
                ['Category', 'Amount'],
                ['Identified Risk Range', f"{self._format_currency(fin.get('total_identified_min'))} - {self._format_currency(fin.get('total_identified_max'))}"],
            ]
            if fin.get('ld_daily_rate'):
                fin_data.append(['Liquidated Damages (Daily)', self._format_currency(fin.get('ld_daily_rate'))])
            if fin.get('ld_cap'):
                fin_data.append(['LD Cap', self._format_currency(fin.get('ld_cap'))])
            if fin.get('bond_percentage'):
                fin_data.append(['Bond Requirement', f"{fin.get('bond_percentage')}%"])
            if fin.get('retention_percentage'):
                fin_data.append(['Retention', f"{fin.get('retention_percentage')}%"])

            fin_table = Table(fin_data, colWidths=[3*inch, 3*inch])
            fin_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1a365d')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
                ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                ('ALIGN', (1, 0), (1, -1), 'RIGHT'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, -1), 10),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
                ('TOPPADDING', (0, 0), (-1, -1), 8),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
            ]))
            elements.append(fin_table)
            elements.append(Spacer(1, 20))

        # Risk Flags Summary
        elements.append(Paragraph("RISK FLAG LOG", self.styles['SectionHeader']))
        elements.append(Paragraph(
            f"Total Flags: {risk_report.total_flags} | High/Critical: {risk_report.high_severity_count}",
            self.styles['Normal']
        ))
        elements.append(Spacer(1, 8))

        # Risk Flags Table
        if risk_report.flags:
            risk_data = [['ID', 'Severity', 'Title', 'Responsibility', 'Est. Cost Impact']]

            for i, flag in enumerate(risk_report.flags):
                if isinstance(flag, dict):
                    risk_id = flag.get('risk_id', f'R{i+1}')
                    severity = flag.get('severity', 'medium').upper()
                    title = flag.get('title', 'Untitled Risk')
                    resp = flag.get('responsibility', 'unknown').upper()
                    cost = flag.get('cost_impact', {})
                    if isinstance(cost, dict) and (cost.get('min_dollars') or cost.get('max_dollars')):
                        cost_str = f"{self._format_currency(cost.get('min_dollars'))} - {self._format_currency(cost.get('max_dollars'))}"
                    elif isinstance(cost, dict) and cost.get('description'):
                        cost_str = cost.get('description', 'TBD')
                    else:
                        cost_str = 'TBD'
                else:
                    # Pydantic model
                    risk_id = getattr(flag, 'risk_id', None) or f'R{i+1}'
                    severity = getattr(flag, 'severity', 'medium')
                    if hasattr(severity, 'value'):
                        severity = severity.value
                    severity = severity.upper()
                    title = getattr(flag, 'title', 'Untitled Risk')
                    resp = getattr(flag, 'responsibility', 'unknown')
                    if hasattr(resp, 'value'):
                        resp = resp.value
                    resp = resp.upper()
                    cost = getattr(flag, 'cost_impact', None)
                    if cost and hasattr(cost, 'min_dollars'):
                        cost_str = f"{self._format_currency(cost.min_dollars)} - {self._format_currency(cost.max_dollars)}"
                    elif cost and hasattr(cost, 'description') and cost.description:
                        cost_str = cost.description
                    else:
                        cost_str = 'TBD'

                # Truncate title if too long
                if len(title) > 40:
                    title = title[:37] + '...'

                risk_data.append([risk_id, severity, title, resp, cost_str])

            risk_table = Table(risk_data, colWidths=[0.5*inch, 0.7*inch, 2.5*inch, 1*inch, 1.5*inch])
            risk_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1a365d')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
                ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, -1), 8),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
                ('TOPPADDING', (0, 0), (-1, -1), 6),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
                ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F9FAFB')]),
            ]))
            elements.append(risk_table)
            elements.append(PageBreak())

        # Detailed Risk Descriptions
        elements.append(Paragraph("DETAILED RISK DESCRIPTIONS", self.styles['SectionHeader']))

        for i, flag in enumerate(risk_report.flags):
            if isinstance(flag, dict):
                risk_id = flag.get('risk_id', f'R{i+1}')
                severity = flag.get('severity', 'medium')
                title = flag.get('title', 'Untitled Risk')
                desc = flag.get('description', '')
                source_quote = flag.get('source_quote', '')
                impact = flag.get('impact', [])
                actions = flag.get('recommended_action', [])
            else:
                risk_id = getattr(flag, 'risk_id', None) or f'R{i+1}'
                severity = getattr(flag, 'severity', 'medium')
                if hasattr(severity, 'value'):
                    severity = severity.value
                title = getattr(flag, 'title', 'Untitled Risk')
                desc = getattr(flag, 'description', '')
                source_quote = getattr(flag, 'source_quote', '') or ''
                impact = getattr(flag, 'impact', []) or []
                actions = getattr(flag, 'recommended_action', []) or []

            sev_color = self._get_severity_color(severity)
            elements.append(Paragraph(
                f"<font color='{sev_color}'>[{severity.upper()}]</font> {risk_id}: {title}",
                self.styles['RiskTitle']
            ))
            if desc:
                elements.append(Paragraph(desc, self.styles['RiskDescription']))
            if source_quote:
                elements.append(Paragraph(
                    f"<i>Source: \"{source_quote}\"</i>",
                    ParagraphStyle('Quote', parent=self.styles['RiskDescription'], textColor=colors.grey)
                ))
            if impact:
                elements.append(Paragraph("<b>Impact:</b>", self.styles['RiskDescription']))
                for item in impact[:3]:
                    elements.append(Paragraph(f"• {item}", self.styles['RiskDescription']))
            if actions:
                elements.append(Paragraph("<b>Recommended Action:</b>", self.styles['RiskDescription']))
                for item in actions[:3]:
                    elements.append(Paragraph(f"• {item}", self.styles['RiskDescription']))
            elements.append(Spacer(1, 8))

        # Estimator Checklist
        if risk_report.estimator_checklist and isinstance(risk_report.estimator_checklist, dict):
            elements.append(PageBreak())
            elements.append(Paragraph("ESTIMATOR CHECKLIST", self.styles['SectionHeader']))

            checklist = risk_report.estimator_checklist

            for section_title, section_key in [
                ("Must Confirm Before Pricing", "must_confirm_before_pricing"),
                ("Include in Bid Cost", "include_in_bid_cost"),
                ("Clarify via RFI", "clarify_via_rfi"),
            ]:
                items = checklist.get(section_key, [])
                if items:
                    elements.append(Paragraph(f"<b>{section_title}</b>", self.styles['Normal']))
                    for item in items:
                        if isinstance(item, dict):
                            item_text = item.get('item', str(item))
                            est_cost = item.get('estimated_cost', '')
                            if est_cost:
                                item_text += f" ({est_cost})"
                        else:
                            item_text = str(item)
                        elements.append(Paragraph(f"☐ {item_text}", self.styles['RiskDescription']))
                    elements.append(Spacer(1, 8))

        # Footer note
        elements.append(Spacer(1, 20))
        elements.append(Paragraph(
            "<i>This report was generated by SpecSentinel AI for estimator documentation purposes. "
            "All identified risks and recommendations should be verified against the original specification documents.</i>",
            ParagraphStyle('Footer', parent=self.styles['Normal'], fontSize=8, textColor=colors.grey)
        ))

        doc.build(elements)
        return buffer.getvalue()

    def generate_excel(
        self,
        risk_report: RiskReport,
        document_name: str,
        project_name: str | None = None,
        extraction: SpecExtraction | None = None,
    ) -> bytes:
        """Generate an Excel risk log."""
        wb = Workbook()

        # Risk Flags Sheet
        ws = wb.active
        ws.title = "Risk Log"

        # Header styling
        header_fill = PatternFill(start_color="1a365d", end_color="1a365d", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        # Title row
        ws.merge_cells('A1:H1')
        ws['A1'] = f"Risk Log - {document_name}"
        ws['A1'].font = Font(size=14, bold=True)

        ws['A2'] = f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
        if project_name:
            ws['A2'] = f"Project: {project_name} | Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}"

        # Headers
        headers = ['ID', 'Severity', 'Title', 'Description', 'Responsibility', 'Est. Min Cost', 'Est. Max Cost', 'Status', 'Notes']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=4, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.border = thin_border

        # Data rows
        severity_fills = {
            'critical': PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid"),
            'high': PatternFill(start_color="FFEDD5", end_color="FFEDD5", fill_type="solid"),
            'medium': PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid"),
            'low': PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid"),
        }

        for i, flag in enumerate(risk_report.flags):
            row = 5 + i

            if isinstance(flag, dict):
                risk_id = flag.get('risk_id', f'R{i+1}')
                severity = flag.get('severity', 'medium')
                title = flag.get('title', '')
                desc = flag.get('description', '')
                resp = flag.get('responsibility', 'unknown')
                cost = flag.get('cost_impact', {})
                min_cost = cost.get('min_dollars') if isinstance(cost, dict) else None
                max_cost = cost.get('max_dollars') if isinstance(cost, dict) else None
            else:
                risk_id = getattr(flag, 'risk_id', None) or f'R{i+1}'
                severity = getattr(flag, 'severity', 'medium')
                if hasattr(severity, 'value'):
                    severity = severity.value
                title = getattr(flag, 'title', '')
                desc = getattr(flag, 'description', '')
                resp = getattr(flag, 'responsibility', 'unknown')
                if hasattr(resp, 'value'):
                    resp = resp.value
                cost = getattr(flag, 'cost_impact', None)
                min_cost = cost.min_dollars if cost and hasattr(cost, 'min_dollars') else None
                max_cost = cost.max_dollars if cost and hasattr(cost, 'max_dollars') else None

            ws.cell(row=row, column=1, value=risk_id).border = thin_border
            sev_cell = ws.cell(row=row, column=2, value=severity.upper())
            sev_cell.border = thin_border
            sev_cell.fill = severity_fills.get(severity.lower(), PatternFill())
            ws.cell(row=row, column=3, value=title).border = thin_border
            ws.cell(row=row, column=4, value=desc).border = thin_border
            ws.cell(row=row, column=5, value=resp.upper()).border = thin_border
            ws.cell(row=row, column=6, value=min_cost).border = thin_border
            ws.cell(row=row, column=7, value=max_cost).border = thin_border
            ws.cell(row=row, column=8, value='OPEN').border = thin_border
            ws.cell(row=row, column=9, value='').border = thin_border

        # Column widths
        ws.column_dimensions['A'].width = 8
        ws.column_dimensions['B'].width = 10
        ws.column_dimensions['C'].width = 35
        ws.column_dimensions['D'].width = 50
        ws.column_dimensions['E'].width = 15
        ws.column_dimensions['F'].width = 12
        ws.column_dimensions['G'].width = 12
        ws.column_dimensions['H'].width = 12
        ws.column_dimensions['I'].width = 30

        # Summary Sheet
        ws_summary = wb.create_sheet("Summary")

        ws_summary['A1'] = "Bid Risk Summary"
        ws_summary['A1'].font = Font(size=14, bold=True)

        summary_data = [
            ('Overall Risk Level', risk_report.overall_risk_level.value if hasattr(risk_report.overall_risk_level, 'value') else str(risk_report.overall_risk_level)),
            ('Total Risk Flags', risk_report.total_flags),
            ('High/Critical Flags', risk_report.high_severity_count),
        ]

        if risk_report.go_no_go and isinstance(risk_report.go_no_go, dict):
            gonogo = risk_report.go_no_go
            summary_data.append(('Go/No-Go Recommendation', gonogo.get('recommendation', '').replace('_', ' ').upper()))
            summary_data.append(('Contingency %', f"{gonogo.get('contingency_percent', 0)}%"))

        if risk_report.financial_exposure and isinstance(risk_report.financial_exposure, dict):
            fin = risk_report.financial_exposure
            summary_data.append(('Min Financial Exposure', fin.get('total_identified_min', 0)))
            summary_data.append(('Max Financial Exposure', fin.get('total_identified_max', 0)))

        for i, (label, value) in enumerate(summary_data):
            ws_summary.cell(row=3+i, column=1, value=label).font = Font(bold=True)
            ws_summary.cell(row=3+i, column=2, value=value)

        ws_summary.column_dimensions['A'].width = 25
        ws_summary.column_dimensions['B'].width = 30

        # Checklist Sheet
        if risk_report.estimator_checklist and isinstance(risk_report.estimator_checklist, dict):
            ws_check = wb.create_sheet("Checklist")
            ws_check['A1'] = "Estimator Checklist"
            ws_check['A1'].font = Font(size=14, bold=True)

            row = 3
            checklist = risk_report.estimator_checklist

            for section_title, section_key in [
                ("Must Confirm Before Pricing", "must_confirm_before_pricing"),
                ("Include in Bid Cost", "include_in_bid_cost"),
                ("Clarify via RFI", "clarify_via_rfi"),
            ]:
                items = checklist.get(section_key, [])
                if items:
                    ws_check.cell(row=row, column=1, value=section_title).font = Font(bold=True)
                    row += 1
                    for item in items:
                        if isinstance(item, dict):
                            item_text = item.get('item', str(item))
                        else:
                            item_text = str(item)
                        ws_check.cell(row=row, column=1, value="☐")
                        ws_check.cell(row=row, column=2, value=item_text)
                        row += 1
                    row += 1

            ws_check.column_dimensions['A'].width = 5
            ws_check.column_dimensions['B'].width = 60

        buffer = io.BytesIO()
        wb.save(buffer)
        return buffer.getvalue()


# Singleton instance
export_service = ExportService()
